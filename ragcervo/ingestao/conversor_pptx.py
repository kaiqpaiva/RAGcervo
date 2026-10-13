import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

TITULOS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
IGNORAR = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}


def tipo_placeholder(shape): # recebe um shape e verifica se é um placeholder
    tipo = None

    if shape.is_placeholder:
        tipo = shape.placeholder_format.type

    return tipo


def textos_da_shape(shape):
    textos = [] # inicia vazio

    if shape.has_text_frame: # só entra se a shape tiver texto
        for paragrafo in shape.text_frame.paragraphs:
            texto = paragrafo.text.strip()
            if texto != "":
                textos.append(texto)

    return textos


def converter_slide(slide, numero):
    titulo = None
    blocos = []

    for shape in slide.shapes:
        tipo = tipo_placeholder(shape)
        textos = textos_da_shape(shape)

        if tipo in IGNORAR:
            continue  # número do slide, rodapé, data

        if len(textos) == 0:
            continue  # imagens, cards e círculos decorativos

        if tipo in TITULOS and titulo is None:
            titulo = " ".join(textos)
        else:
            blocos.append(textos)

    # plano B: sem placeholder de título com texto, o primeiro texto vira título
    if titulo is None and len(blocos) > 0:
        primeiro_bloco = blocos.pop(0)
        titulo = " ".join(primeiro_bloco)

    if titulo is None:
        cabecalho = f"## Slide {numero}"
    else:
        cabecalho = f"## Slide {numero}: {titulo}"

    linhas = [cabecalho]
    for textos in blocos:
        for texto in textos:
            linhas.append(texto)

    return "\n\n".join(linhas)


def converter_arquivo(caminho):
    caminho = Path(caminho)
    apresentacao = Presentation(caminho)

    partes = [f"# {caminho.stem}"]
    numero = 1
    for slide in apresentacao.slides:
        partes.append(converter_slide(slide, numero))
        numero = numero + 1

    return "\n\n".join(partes) + "\n"


if __name__ == "__main__":
    origem = Path(sys.argv[1])
    destino = Path("dados") / "markdown" / f"{origem.stem}.md"

    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(converter_arquivo(origem), encoding="utf-8")

    print(f"salvo em {destino}")