import sys
from pptx import Presentation

prs = Presentation(sys.argv[1])

for i, slide in enumerate(prs.slides, start=1):
    print(f"Slide {i}")
    for shape in slide.shapes:
        tipo_ph = shape.placeholder_format.type if shape.is_placeholder else ""
        print(f"\t{shape.shape_type} | {shape.name} | {tipo_ph}")
        if shape.has_text_frame:
            for par in shape.text_frame.paragraphs:
                print("\t\t" + "\t" * par.level + f"[nível {par.level}] {par.text}")
    if slide.has_notes_slide:
        print(f"\tNotas: {slide.notes_slide.notes_text_frame.text}")